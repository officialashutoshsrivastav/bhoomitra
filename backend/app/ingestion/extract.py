from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from app.core.config import settings
from app.core.logging import get_logger

logger = get_logger("bhoomitra.ingestion")


@dataclass
class PageText:
    page: int
    text: str
    ocr: bool = False


@dataclass
class Extraction:
    pages: List[PageText]
    metadata: Dict[str, object] = field(default_factory=dict)
    ocr_applied: bool = False

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


# ---------------------------------------------------------------------------
# OCR providers
# ---------------------------------------------------------------------------
class OCRProvider(ABC):
    name = "abstract"

    @abstractmethod
    def available(self) -> bool: ...

    @abstractmethod
    def image_to_text(self, image_bytes: bytes) -> str: ...


class TesseractOCR(OCRProvider):
    name = "tesseract"

    def available(self) -> bool:
        return shutil.which("tesseract") is not None

    def image_to_text(self, image_bytes: bytes) -> str:
        import pytesseract
        from PIL import Image

        img = Image.open(BytesIO(image_bytes))
        return pytesseract.image_to_string(img, lang="eng")


class NullOCR(OCRProvider):
    name = "none"

    def available(self) -> bool:
        return False

    def image_to_text(self, image_bytes: bytes) -> str:
        return ""


def get_ocr_provider() -> OCRProvider:
    t = TesseractOCR()
    return t if (settings.ocr_enabled and t.available()) else NullOCR()


def _render_pdf_pages(pdf_bytes: bytes, pages: List[int], dpi: int = 200) -> Dict[int, bytes]:
    """Rasterise selected pages with poppler's pdftoppm (if installed)."""
    if shutil.which("pdftoppm") is None:
        return {}
    out: Dict[int, bytes] = {}
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "doc.pdf"
        src.write_bytes(pdf_bytes)
        for p in pages:
            prefix = Path(tmp) / f"p{p}"
            try:
                subprocess.run(["pdftoppm", "-r", str(dpi), "-f", str(p), "-l", str(p), "-png", str(src), str(prefix)], check=True, capture_output=True, timeout=120)
            except Exception as exc:
                logger.warning("pdftoppm failed", extra={"extra_data": {"page": p, "error": str(exc)[:200]}})
                continue
            for f in Path(tmp).glob(f"p{p}*.png"):
                out[p] = f.read_bytes()
                break
    return out


# ---------------------------------------------------------------------------
# Extractors
# ---------------------------------------------------------------------------
def extract_pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    meta: Dict[str, object] = {"page_count": len(reader.pages)}
    try:
        info = reader.metadata or {}
        for k in ("/Title", "/Author", "/Subject", "/Producer", "/CreationDate"):
            if info.get(k):
                meta[k.strip("/").lower()] = str(info.get(k))
    except Exception:
        pass
    pages: List[PageText] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        pages.append(PageText(page=i, text=clean_text(txt)))

    ocr = get_ocr_provider()
    weak = [p.page for p in pages if len(p.text) < settings.ocr_min_chars_per_page]
    ocr_applied = False
    if weak and ocr.available():
        images = _render_pdf_pages(data, weak[:50])
        for p in pages:
            if p.page in images:
                try:
                    txt = clean_text(ocr.image_to_text(images[p.page]))
                except Exception as exc:
                    logger.warning("ocr failed", extra={"extra_data": {"page": p.page, "error": str(exc)[:200]}})
                    continue
                if len(txt) > len(p.text):
                    p.text, p.ocr, ocr_applied = txt, True, True
    meta["ocr_pages"] = [p.page for p in pages if p.ocr]
    meta["ocr_provider"] = ocr.name
    return Extraction(pages=pages, metadata=meta, ocr_applied=ocr_applied)


def extract_docx(data: bytes) -> Extraction:
    import docx

    d = docx.Document(BytesIO(data))
    paras = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    text = clean_text("\n".join(paras))
    meta = {"paragraphs": len(paras)}
    try:
        cp = d.core_properties
        if cp.title:
            meta["title"] = cp.title
        if cp.author:
            meta["author"] = cp.author
    except Exception:
        pass
    return Extraction(pages=_paginate_plain(text), metadata=meta)


def extract_image(data: bytes) -> Extraction:
    ocr = get_ocr_provider()
    if not ocr.available():
        return Extraction(pages=[PageText(page=1, text="", ocr=False)], metadata={"ocr_provider": "unavailable"}, ocr_applied=False)
    text = clean_text(ocr.image_to_text(data))
    return Extraction(pages=[PageText(page=1, text=text, ocr=True)], metadata={"ocr_provider": ocr.name}, ocr_applied=True)


def extract_text_file(data: bytes) -> Extraction:
    text = clean_text(data.decode("utf-8", errors="replace"))
    return Extraction(pages=_paginate_plain(text), metadata={})


def _paginate_plain(text: str, chars_per_page: int = 3500) -> List[PageText]:
    if not text:
        return [PageText(page=1, text="")]
    pages, buf, n = [], [], 0
    for para in text.split("\n\n"):
        buf.append(para)
        n += len(para)
        if n >= chars_per_page:
            pages.append(PageText(page=len(pages) + 1, text="\n\n".join(buf)))
            buf, n = [], 0
    if buf:
        pages.append(PageText(page=len(pages) + 1, text="\n\n".join(buf)))
    return pages


def extract(data: bytes, mime: str) -> Extraction:
    if mime == "application/pdf":
        return extract_pdf(data)
    if mime.endswith("wordprocessingml.document"):
        return extract_docx(data)
    if mime.startswith("image/"):
        return extract_image(data)
    if mime in {"text/plain", "text/markdown", "text/csv", "application/json", "application/geo+json"}:
        return extract_text_file(data)
    return Extraction(pages=[PageText(page=1, text="")], metadata={"note": f"no text extractor for {mime}"})


# ---------------------------------------------------------------------------
# Cleaning + chunking
# ---------------------------------------------------------------------------
_WS = re.compile(r"[ \t\u00a0]+")
_NL = re.compile(r"\n{3,}")
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    text = _WS.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = _NL.sub("\n\n", text)
    return text.strip()


_SENT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9(\"'])")


@dataclass
class Chunk:
    index: int
    content: str
    page_start: int
    page_end: int


def chunk_pages(pages: List[PageText], size: Optional[int] = None, overlap: Optional[int] = None) -> List[Chunk]:
    size = size or settings.chunk_size_chars
    overlap = overlap or settings.chunk_overlap_chars
    units: List[Tuple[str, int]] = []
    for p in pages:
        if not p.text:
            continue
        for para in p.text.split("\n\n"):
            para = para.strip()
            if not para:
                continue
            if len(para) > size:
                for sent in _SENT.split(para):
                    if sent.strip():
                        units.append((sent.strip(), p.page))
            else:
                units.append((para, p.page))

    chunks: List[Chunk] = []
    buf: List[Tuple[str, int]] = []
    cur = 0

    def flush():
        nonlocal buf, cur
        if not buf:
            return
        text = "\n".join(u for u, _ in buf).strip()
        if text:
            chunks.append(Chunk(index=len(chunks), content=text, page_start=buf[0][1], page_end=buf[-1][1]))
        # keep tail for overlap
        keep: List[Tuple[str, int]] = []
        acc = 0
        for u in reversed(buf):
            if acc + len(u[0]) > overlap:
                break
            keep.insert(0, u)
            acc += len(u[0])
        buf = keep
        cur = acc

    for unit, page in units:
        if cur + len(unit) > size and buf:
            flush()
        if len(unit) > size * 2:
            unit = unit[: size * 2]
        buf.append((unit, page))
        cur += len(unit)
    flush()
    # de-duplicate degenerate overlap-only chunks
    out: List[Chunk] = []
    seen = set()
    for c in chunks:
        key = c.content[:200]
        if key in seen:
            continue
        seen.add(key)
        c.index = len(out)
        out.append(c)
    return out


def estimate_tokens(text: str) -> int:
    return max(1, len(text) // 4)
